"""Project an immutable :mod:`layout_model` deck into an editable PPTX file.

This module is intentionally a serializer.  It consumes resolved rectangles,
typography, list indents, crops, and reveal targets; physical-layout choices
belong to ``layout_engine`` and are not repeated here.
"""

from __future__ import annotations

import os
import pathlib
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

import slide_lib.layout_content
import slide_lib.layout_model
import slide_lib.layout_primitives
import slide_lib.pptx_animation
import slide_lib.presentation_theme


_ROLE_COLORS = {
	slide_lib.layout_primitives.StyleRole.ACCENT: "24578F",
	slide_lib.layout_primitives.StyleRole.TABLE_HEADER: "24578F",
	slide_lib.layout_primitives.StyleRole.MUTED: "526176",
	slide_lib.layout_primitives.StyleRole.DECORATION: "FFFFFF",
}
_FOREGROUND = "172033"
_WHITE = "FFFFFF"


def write_pptx(deck: slide_lib.layout_model.LayoutDeck,
		theme: slide_lib.presentation_theme.PresentationTheme,
		destination: pathlib.Path) -> pathlib.Path:
	"""Atomically write ``deck`` as editable PPTX using only resolved plan facts."""
	if not isinstance(deck, slide_lib.layout_model.LayoutDeck):
		raise TypeError("write_pptx requires a LayoutDeck")
	if not isinstance(theme, slide_lib.presentation_theme.PresentationTheme):
		raise TypeError("write_pptx requires a PresentationTheme")
	destination = pathlib.Path(destination)
	destination.parent.mkdir(parents=True, exist_ok=True)
	presentation = Presentation()
	presentation.slide_width = _emu(deck.canvas.width, theme)
	presentation.slide_height = _emu(deck.canvas.height, theme)
	title = next((item.value for item in deck.metadata if item.name == "title"), deck.identity.deck_id)
	presentation.core_properties.title = title
	blank = presentation.slide_layouts[6]
	for plan_slide in deck.slides:
		slide = presentation.slides.add_slide(blank)
		_write_slide(slide, plan_slide, deck, theme)
	with tempfile.NamedTemporaryFile(prefix=f".{destination.stem}.", suffix=".pptx",
		dir=destination.parent, delete=False) as temporary:
		temporary_path = pathlib.Path(temporary.name)
	try:
		presentation.save(temporary_path)
		os.replace(temporary_path, destination)
	except BaseException:
		temporary_path.unlink(missing_ok=True)
		raise
	return destination


def _write_slide(slide: object, plan: slide_lib.layout_model.LayoutSlide,
		deck: slide_lib.layout_model.LayoutDeck,
		theme: slide_lib.presentation_theme.PresentationTheme) -> None:
	"""Write one resolved physical slide in plan z/reading order."""
	_write_background(slide, deck, theme)
	writer = slide_lib.pptx_animation.PptxAnimationWriter(slide)
	shapes: dict[str, object] = {}
	for item in sorted(plan.objects, key=lambda value: (value.z_index, value.reading_order, value.object_id)):
		shapes[item.object_id] = _write_object(slide, item, deck, theme, plan.continuation_context)
	targets = sorted((target for item in plan.objects for target in item.reveal_targets),
		key=lambda value: value.activation_order)
	for target in targets:
		writer.register(shapes[target.object_id], target.reveal,
			(target.paragraph_indexes[0], target.paragraph_indexes[-1])
			if target.paragraph_indexes else None)
	writer.finalize()
	_write_notes(slide, plan)


def _write_background(slide: object, deck: slide_lib.layout_model.LayoutDeck,
		theme: slide_lib.presentation_theme.PresentationTheme) -> None:
	"""Project the supplied template background without deriving layout geometry."""
	slide.background.fill.solid()
	slide.background.fill.fore_color.rgb = RGBColor.from_string(_WHITE)
	band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(0, theme), _emu(0, theme),
		_emu(deck.canvas.width, theme), _emu(theme.top_band_height, theme))
	band.fill.solid()
	band.fill.fore_color.rgb = RGBColor.from_string(theme.gradient_start_color)
	_apply_top_band_gradient(band, theme)
	band.line.fill.background()
	_set_accessibility(band, "Theme background", "Lecture template top band", False)


def _apply_top_band_gradient(shape: object,
		theme: slide_lib.presentation_theme.PresentationTheme) -> None:
	"""Replace one solid fill with the validated lecture-theme gradient."""
	properties = shape.element.spPr
	solid_fill = next(child for child in properties if child.tag.endswith("solidFill"))
	gradient = OxmlElement("a:gradFill")
	stops = OxmlElement("a:gsLst")
	# ASVS 2.2.1: gradient colors come only from the validated template contract.
	for position, color in ((0, theme.gradient_start_color),
			(100000, theme.gradient_end_color)):
		stop = OxmlElement("a:gs")
		stop.set("pos", str(position))
		value = OxmlElement("a:srgbClr")
		value.set("val", color)
		stop.append(value)
		stops.append(stop)
	gradient.append(stops)
	linear = OxmlElement("a:lin")
	linear.set("ang", "5400000")
	linear.set("scaled", "1")
	gradient.append(linear)
	properties.replace(solid_fill, gradient)


def _write_object(slide: object, item: slide_lib.layout_model.LayoutObject,
		deck: slide_lib.layout_model.LayoutDeck,
		theme: slide_lib.presentation_theme.PresentationTheme,
		continuation_context: slide_lib.layout_model.ContinuationContext | None) -> object:
	"""Write one object using its already-selected representation and rectangle."""
	content = item.content
	if isinstance(content, slide_lib.layout_content.TextContent):
		shape = _add_textbox(slide, item.rectangle, item.frame_text, theme)
		_write_text(shape.text_frame, content, theme)
		_set_accessibility(shape, *_continuation_accessibility(item, _text_accessibility(item), continuation_context))
		return shape
	if isinstance(content, slide_lib.layout_content.TableContent):
		shape = _write_table(slide, item.rectangle, content, theme)
		_set_accessibility(shape, *_continuation_accessibility(item, _text_accessibility(item), continuation_context))
		return shape
	if isinstance(content, slide_lib.layout_content.PictureContent):
		shape = _write_picture(slide, content, deck, theme)
		_set_accessibility(shape, *_continuation_accessibility(item,
			(content.accessibility.name, content.accessibility.description, content.accessibility.decorative), continuation_context))
		return shape
	if isinstance(content, slide_lib.layout_content.ShapeContent):
		shape = _write_shape(slide, item.rectangle, content, item.frame_text, theme)
		_set_accessibility(shape, *_continuation_accessibility(item,
			(content.accessibility.name, content.accessibility.description, content.accessibility.decorative), continuation_context))
		return shape
	raise ValueError(f"unsupported planned PPTX content: {type(content).__name__}")


def _add_textbox(slide: object, rectangle: slide_lib.layout_primitives.LogicalRectangle,
		frame: slide_lib.layout_primitives.FrameTextProperties | None,
		theme: slide_lib.presentation_theme.PresentationTheme) -> object:
	shape = slide.shapes.add_textbox(_emu(rectangle.x, theme), _emu(rectangle.y, theme),
		_emu(rectangle.width, theme), _emu(rectangle.height, theme))
	text_frame = shape.text_frame
	text_frame.clear()
	if frame is not None:
		text_frame.margin_left = _emu(frame.padding.left, theme)
		text_frame.margin_right = _emu(frame.padding.right, theme)
		text_frame.margin_top = _emu(frame.padding.top, theme)
		text_frame.margin_bottom = _emu(frame.padding.bottom, theme)
		text_frame.vertical_anchor = _vertical_anchor(frame.vertical_alignment)
		text_frame.word_wrap = frame.wrap is slide_lib.layout_primitives.TextWrap.WRAP
		body = text_frame._txBody.bodyPr
		if frame.text_direction is slide_lib.layout_primitives.TextDirection.VERTICAL:
			body.set("vert", "vert")
		if frame.overflow_policy is slide_lib.layout_primitives.OverflowPolicy.SHRINK:
			for child in list(body):
				if child.tag.endswith(("noAutofit", "normAutofit", "spAutoFit")):
					body.remove(child)
			autofit = OxmlElement("a:normAutofit")
			autofit.set("fontScale", "100000")
			autofit.set("lnSpcReduction", "0")
			body.append(autofit)
	return shape


def _write_text(frame: object, content: slide_lib.layout_content.TextContent,
		theme: slide_lib.presentation_theme.PresentationTheme) -> None:
	for index, paragraph_data in enumerate(content.paragraphs):
		paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
		_apply_paragraph(paragraph, paragraph_data, theme)
		for inline in paragraph_data.inlines:
			if isinstance(inline, slide_lib.layout_content.LineBreak):
				paragraph.add_line_break()
			else:
				_write_run(paragraph.add_run(), inline, paragraph_data.typography)
		if not paragraph.runs and not paragraph._p.findall(".//a:br", {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}):
			_write_run(paragraph.add_run(), slide_lib.layout_content.TextRun("", _fallback_style()), paragraph_data.typography)


def _apply_paragraph(paragraph: object, data: slide_lib.layout_content.TextParagraph,
		theme: slide_lib.presentation_theme.PresentationTheme) -> None:
	properties = data.properties
	paragraph.alignment = _paragraph_alignment(properties.horizontal_alignment)
	paragraph.space_before = Pt(properties.space_before_pt)
	paragraph.space_after = Pt(properties.space_after_pt)
	paragraph.line_spacing = Pt(properties.line_spacing_pt)
	ppr = paragraph._p.get_or_add_pPr()
	ppr.set("marL", str(int(round(properties.indent_start_pt * 12700))))
	ppr.set("marR", str(int(round(properties.indent_end_pt * 12700))))
	ppr.set("indent", str(int(round(properties.first_line_indent_pt * 12700))))
	for child in list(ppr):
		if child.tag.endswith(("buNone", "buAutoNum", "buChar", "tabLst")):
			ppr.remove(child)
	if data.list_metadata is None:
		ppr.append(OxmlElement("a:buNone"))
	else:
		metadata = data.list_metadata
		paragraph.level = metadata.level
		bullet = OxmlElement("a:buAutoNum" if metadata.kind is slide_lib.layout_primitives.ListKind.ORDERED else "a:buChar")
		if metadata.kind is slide_lib.layout_primitives.ListKind.ORDERED:
			bullet.set("type", "arabicPeriod")
			bullet.set("startAt", str(metadata.start))
		else:
			bullet.set("char", _bullet_character(theme, metadata.level))
		ppr.append(bullet)
	if properties.tab_stops_pt:
		tabs = OxmlElement("a:tabLst")
		for value in properties.tab_stops_pt:
			tabs.append(OxmlElement("a:tab", pos=str(int(round(value * 12700)))))
		ppr.append(tabs)


def _write_run(run: object, source: slide_lib.layout_content.TextRun,
		typography: slide_lib.layout_primitives.Typography) -> None:
	style = source.style
	run.text = source.text
	run.font.name = style.font_family
	run.font.size = Pt(typography.selected_size_pt)
	run.font.bold = style.bold
	run.font.italic = style.italic
	run.font.underline = style.underline
	run.font.color.rgb = RGBColor.from_string(style.foreground)
	if style.link_url is not None:
		run.hyperlink.address = style.link_url


def _write_table(slide: object, rectangle: slide_lib.layout_primitives.LogicalRectangle,
		content: slide_lib.layout_content.TableContent,
		theme: slide_lib.presentation_theme.PresentationTheme) -> object:
	rows = content.header_rows + content.body_rows
	shape = slide.shapes.add_table(len(rows), len(content.column_widths), _emu(rectangle.x, theme),
		_emu(rectangle.y, theme), _emu(rectangle.width, theme), _emu(rectangle.height, theme))
	table = shape.table
	for index, width in enumerate(content.column_widths):
		table.columns[index].width = _emu(width, theme)
	for index, height in enumerate(content.row_heights):
		table.rows[index].height = _emu(height, theme)
	occupied = [[False] * len(content.column_widths) for _ in rows]
	for row_index, row in enumerate(rows):
		column = 0
		for cell_data in row.cells:
			while occupied[row_index][column]:
				column += 1
			cell = table.cell(row_index, column)
			cell.margin_left = _emu(cell_data.padding.left, theme)
			cell.margin_right = _emu(cell_data.padding.right, theme)
			cell.margin_top = _emu(cell_data.padding.top, theme)
			cell.margin_bottom = _emu(cell_data.padding.bottom, theme)
			cell.vertical_anchor = _vertical_anchor(cell_data.vertical_alignment)
			if row_index < len(content.header_rows):
				cell.fill.solid()
				cell.fill.fore_color.rgb = RGBColor.from_string(_color_for(content.style.header_role))
			_write_text(cell.text_frame, slide_lib.layout_content.TextContent(cell_data.paragraphs), theme)
			for y in range(row_index, row_index + cell_data.row_span):
				for x in range(column, column + cell_data.column_span):
					occupied[y][x] = True
			if cell_data.column_span > 1 or cell_data.row_span > 1:
				cell.merge(table.cell(row_index + cell_data.row_span - 1, column + cell_data.column_span - 1))
			column += cell_data.column_span
	return shape


def _write_picture(slide: object, content: slide_lib.layout_content.PictureContent,
		deck: slide_lib.layout_model.LayoutDeck,
		theme: slide_lib.presentation_theme.PresentationTheme) -> object:
	path = pathlib.Path(content.source_path)
	if not path.is_absolute():
		path = deck.identity.source_path.parent / path
	if not path.is_file():
		raise ValueError(f"planned picture source is missing: {content.source_path}")
	placement = content.placement
	display = placement.displayed_rectangle
	shape = slide.shapes.add_picture(str(path), _emu(display.x, theme), _emu(display.y, theme),
		_emu(display.width, theme), _emu(display.height, theme))
	shape.crop_left = placement.crop.left
	shape.crop_top = placement.crop.top
	shape.crop_right = placement.crop.right
	shape.crop_bottom = placement.crop.bottom
	if content.link_url is not None:
		shape.click_action.hyperlink.address = content.link_url
	return shape


def _write_shape(slide: object, rectangle: slide_lib.layout_primitives.LogicalRectangle,
		content: slide_lib.layout_content.ShapeContent,
		frame: slide_lib.layout_primitives.FrameTextProperties | None,
		theme: slide_lib.presentation_theme.PresentationTheme) -> object:
	if content.kind is slide_lib.layout_primitives.ShapeKind.LINE:
		shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
			_emu(rectangle.x, theme), _emu(rectangle.y, theme),
			_emu(rectangle.x + rectangle.width, theme),
			_emu(rectangle.y + rectangle.height, theme))
	else:
		kind = {slide_lib.layout_primitives.ShapeKind.RECTANGLE: MSO_SHAPE.RECTANGLE,
			slide_lib.layout_primitives.ShapeKind.ROUNDED_RECTANGLE:
			MSO_SHAPE.ROUNDED_RECTANGLE}[content.kind]
		shape = slide.shapes.add_shape(kind, _emu(rectangle.x, theme),
			_emu(rectangle.y, theme), _emu(rectangle.width, theme),
			_emu(rectangle.height, theme))
		shape.fill.solid()
		shape.fill.fore_color.rgb = RGBColor.from_string(
			_color_for(content.style.fill_role))
	shape.line.color.rgb = RGBColor.from_string(_color_for(content.style.line_role))
	shape.line.width = Pt(content.style.line_width_pt)
	if content.text is not None:
		text_frame = shape.text_frame
		text_frame.clear()
		if frame is not None:
			text_frame.margin_left = _emu(frame.padding.left, theme)
			text_frame.margin_right = _emu(frame.padding.right, theme)
			text_frame.margin_top = _emu(frame.padding.top, theme)
			text_frame.margin_bottom = _emu(frame.padding.bottom, theme)
			text_frame.vertical_anchor = _vertical_anchor(frame.vertical_alignment)
		_write_text(text_frame, content.text, theme)
	if content.link_url is not None:
		shape.click_action.hyperlink.address = content.link_url
	return shape


def _write_notes(slide: object, plan: slide_lib.layout_model.LayoutSlide) -> None:
	notes = [item.text for item in plan.notes]
	if plan.continuation_context is not None:
		trail = " > ".join(_inline_text(item.inlines) for item in plan.continuation_context.entries)
		notes.append(f"Continuation context ({plan.continuation_context.display.value}): {trail}")
	if notes:
		slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)


def _continuation_accessibility(item: slide_lib.layout_model.LayoutObject,
		base: tuple[str | None, str | None, bool],
		context: slide_lib.layout_model.ContinuationContext | None) -> tuple[str | None, str | None, bool]:
	name, description, decorative = base
	trails = []
	if item.decomposition_origin is not None:
		origin = item.decomposition_origin
		trails.append(f"Decomposed from {origin.source_slide_id}, {origin.original_layout}/{origin.original_slot}.")
	if context is not None:
		trail = " > ".join(_inline_text(entry.inlines) for entry in context.entries)
		trails.append(f"Continuation context ({context.display.value}): {trail}.")
	if not trails:
		return base
	addition = " ".join(trails)
	return name, f"{description} {addition}" if description else addition, decorative


def _text_accessibility(item: slide_lib.layout_model.LayoutObject) -> tuple[str | None, str | None, bool]:
	return item.object_id, _inline_text(tuple(
		inline for paragraph in _text_paragraphs(item.content) for inline in paragraph.inlines)), False


def _text_paragraphs(content: object) -> tuple[slide_lib.layout_content.TextParagraph, ...]:
	if isinstance(content, slide_lib.layout_content.TextContent):
		return content.paragraphs
	if isinstance(content, slide_lib.layout_content.TableContent):
		return tuple(paragraph for row in content.header_rows + content.body_rows for cell in row.cells
			for paragraph in cell.paragraphs)
	return ()


def _inline_text(inlines: tuple[slide_lib.layout_content.InlineContent, ...]) -> str:
	return "".join(item.text if isinstance(item, slide_lib.layout_content.TextRun) else "\n" for item in inlines)


def _set_accessibility(shape: object, name: str | None, description: str | None, decorative: bool) -> None:
	nonvisual = shape.element.xpath(".//p:cNvPr")
	if len(nonvisual) != 1:
		raise ValueError("PPTX object has no unique nonvisual properties")
	properties = nonvisual[0]
	if decorative:
		properties.set("descr", "")
		return
	if name is not None:
		properties.set("name", name)
	if description is not None:
		properties.set("descr", description)


def _emu(value: float, theme: slide_lib.presentation_theme.PresentationTheme) -> Emu:
	return Emu(round(value * theme.emu_per_logical_pixel))


def _color_for(role: slide_lib.layout_primitives.StyleRole) -> str:
	return _ROLE_COLORS.get(role, _FOREGROUND)


def _fallback_style() -> slide_lib.layout_content.RunStyle:
	return slide_lib.layout_content.RunStyle("OpenDyslexic", _FOREGROUND)


def _paragraph_alignment(value: slide_lib.layout_primitives.HorizontalAlignment) -> PP_ALIGN:
	return {slide_lib.layout_primitives.HorizontalAlignment.START: PP_ALIGN.LEFT,
		slide_lib.layout_primitives.HorizontalAlignment.CENTER: PP_ALIGN.CENTER,
		slide_lib.layout_primitives.HorizontalAlignment.END: PP_ALIGN.RIGHT,
		slide_lib.layout_primitives.HorizontalAlignment.JUSTIFY: PP_ALIGN.JUSTIFY}[value]


def _vertical_anchor(value: slide_lib.layout_primitives.VerticalAlignment) -> MSO_ANCHOR:
	return {slide_lib.layout_primitives.VerticalAlignment.TOP: MSO_ANCHOR.TOP,
		slide_lib.layout_primitives.VerticalAlignment.MIDDLE: MSO_ANCHOR.MIDDLE,
		slide_lib.layout_primitives.VerticalAlignment.BOTTOM: MSO_ANCHOR.BOTTOM}[value]


def _bullet_character(theme: slide_lib.presentation_theme.PresentationTheme, level: int) -> str:
	if level >= len(theme.list_levels):
		raise ValueError(f"planned list nesting exceeds theme level count: {level}")
	return theme.list_levels[level].bullet_character

"""Read bounded PPTX containers into raw imported-slide facts."""

# Standard Library
import io
import hashlib
import pathlib
import re
import zipfile

# PIP3 modules
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

# local repo modules
import slide_lib.importers.source_model as source_model


MAX_INPUT_BYTES = 256 * 1024 * 1024
MAX_MEMBER_BYTES = 128 * 1024 * 1024
MAX_UNPACKED_BYTES = 512 * 1024 * 1024
# OOXML presentations commonly contain more relationship parts than ODP files.
MAX_ARCHIVE_MEMBERS = 4000
MAX_IMAGE_PIXELS = 100_000_000
SUPPORTED_IMAGE_SUFFIXES = {".emf", ".gif", ".jpg", ".png", ".wmf"}
SAFE_LINK_SCHEMES = ("http://", "https://", "mailto:")
WMF_HEADERS = (b"\xd7\xcd\xc6\x9a", b"\x01\x00\x09\x00\x00\x03", b"\x02\x00\x09\x00\x00\x03")
EMF_SIGNATURE_OFFSET = 40
EMF_SIGNATURE = b" EMF"


#============================================
def validate_member_name(member_name: str) -> None:
	"""Reject absolute and traversal paths in an OOXML archive."""
	# ASVS 5.3.3: archive member paths never control filesystem destinations.
	normalized = member_name.replace("\\", "/")
	parts = pathlib.PurePosixPath(normalized).parts
	if normalized.startswith("/") or ".." in parts:
		raise ValueError(f"unsafe archive member path: {member_name}")


#============================================
def validate_pptx(input_path: pathlib.Path) -> None:
	"""Validate a bounded OOXML presentation before parsing it."""
	# ASVS 1.5.2, 2.2.1, 5.2.1, 5.2.2, and 5.2.3: validate the type and archive limits.
	if not input_path.is_file() or input_path.suffix.lower() != ".pptx":
		raise ValueError("input must be an existing .pptx file")
	if input_path.stat().st_size > MAX_INPUT_BYTES:
		raise ValueError("PPTX exceeds the compressed input limit")
	with zipfile.ZipFile(input_path) as archive:
		members = archive.infolist()
		if len(members) > MAX_ARCHIVE_MEMBERS:
			raise ValueError("PPTX contains too many archive members")
		total_size = 0
		member_names: set[str] = set()
		for member in members:
			validate_member_name(member.filename)
			if member.file_size > MAX_MEMBER_BYTES:
				raise ValueError(f"PPTX member exceeds size limit: {member.filename}")
			total_size += member.file_size
			if total_size > MAX_UNPACKED_BYTES:
				raise ValueError("PPTX exceeds the expanded archive limit")
			member_names.add(member.filename)
		required_names = {"[Content_Types].xml", "ppt/presentation.xml"}
		if not required_names.issubset(member_names):
			raise ValueError("PPTX is missing required OOXML presentation members")


#============================================
def safe_hyperlink(address: str | None) -> str:
	"""Return an allowed raw hyperlink address or an empty string."""
	if address is None:
		return ""
	# ASVS 1.2.2: only known-safe URL schemes enter generated Djot links.
	if not address.lower().startswith(SAFE_LINK_SCHEMES):
		return ""
	return address


#============================================
def trim_runs(runs: list[source_model.TextRun]) -> tuple[source_model.TextRun, ...]:
	"""Strip paragraph-edge whitespace without changing internal run ownership."""
	while runs and not runs[0].text:
		runs.pop(0)
	while runs and not runs[-1].text:
		runs.pop()
	if not runs:
		return ()
	runs[0] = source_model.TextRun(runs[0].text.lstrip(), runs[0].link)
	runs[-1] = source_model.TextRun(runs[-1].text.rstrip(), runs[-1].link)
	return tuple(run for run in runs if run.text)


#============================================
def paragraph_runs(paragraph: object) -> tuple[source_model.TextRun, ...]:
	"""Read visible paragraph characters and safe links without output escaping."""
	runs: list[source_model.TextRun] = []
	for run in paragraph.runs:
		text = run.text
		if not text:
			continue
		runs.append(source_model.TextRun(text, safe_hyperlink(run.hyperlink.address)))
	if runs:
		return trim_runs(runs)
	text = paragraph.text.strip()
	if not text:
		return ()
	return (source_model.TextRun(text),)


#============================================
def plain_text(runs: tuple[source_model.TextRun, ...]) -> str:
	"""Return visible characters for reader-side density and identity checks."""
	return "".join(run.text for run in runs)


#============================================
def cell_runs(cell: object) -> tuple[source_model.TextRun, ...]:
	"""Read one table cell into the inline vocabulary supported by native tables."""
	result: list[source_model.TextRun] = []
	for paragraph in cell.text_frame.paragraphs:
		paragraph_content = paragraph_runs(paragraph)
		if not paragraph_content:
			continue
		if result:
			result.append(source_model.TextRun(" "))
		result.extend(paragraph_content)
	return tuple(result)


#============================================
def title_shape_id(slide: object) -> int | None:
	"""Return the title shape id without relying on proxy identity."""
	shape = slide.shapes.title
	if shape is None:
		return None
	return shape.shape_id


#============================================
def is_subtitle_shape(shape: object) -> bool:
	"""Return whether a placeholder is a subtitle."""
	if not shape.is_placeholder:
		return False
	placeholder_name = str(shape.placeholder_format.type).upper()
	return "SUBTITLE" in placeholder_name


#============================================
def validate_image_blob(blob: bytes, suffix: str) -> None:
	"""Validate one raster, WMF, or EMF image before writing it."""
	# ASVS 5.2.2 and 5.2.6: validate the declared type and bounded content.
	if suffix not in SUPPORTED_IMAGE_SUFFIXES:
		raise ValueError(f"unsupported PPTX image type: {suffix}")
	if len(blob) > MAX_MEMBER_BYTES:
		raise ValueError("PPTX image exceeds the per-image size limit")
	if suffix == ".wmf":
		if not blob.startswith(WMF_HEADERS):
			raise ValueError("invalid WMF image header")
		return
	if suffix == ".emf":
		if blob[:4] != b"\x01\x00\x00\x00" or blob[EMF_SIGNATURE_OFFSET:44] != EMF_SIGNATURE:
			raise ValueError("invalid EMF image header")
		return
	with Image.open(io.BytesIO(blob)) as image:
		width, height = image.size
		if width < 1 or height < 1 or width * height > MAX_IMAGE_PIXELS:
			raise ValueError("PPTX image dimensions exceed the supported limit")
		image.verify()


#============================================
def image_suffix(blob: bytes, declared_suffix: str) -> str:
	"""Correct LibreOffice's EMF-as-WMF metadata before validation."""
	if declared_suffix == ".wmf" and blob[EMF_SIGNATURE_OFFSET:44] == EMF_SIGNATURE:
		return ".emf"
	return declared_suffix


#============================================
def image_asset(
	shape: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	known_images: dict[str, str],
) -> source_model.ImageAsset:
	"""Extract one picture into an internally named content asset."""
	blob = shape.image.blob
	suffix = f".{shape.image.ext.lower()}"
	if suffix == ".jpeg":
		suffix = ".jpg"
	suffix = image_suffix(blob, suffix)
	validate_image_blob(blob, suffix)
	digest = hashlib.sha256(blob).hexdigest()
	asset_name = known_images.get(digest)
	if asset_name is None:
		asset_name = f"image_{len(known_images) + 1:03d}{suffix}"
		# ASVS 5.3.2: archive metadata never selects the output destination.
		(assets_dir / asset_name).write_bytes(blob)
		known_images[digest] = asset_name
	asset_path = (djot_root / asset_name).as_posix()
	asset_number = int(re.search(r"\d+", asset_name).group())
	return source_model.ImageAsset(
		shape.left, shape.top, shape.width, shape.height, asset_path,
		f"Slide image {asset_number}",
	)


#============================================
def table_block(shape: object) -> source_model.TableBlock:
	"""Read one actual source table without flattening its rows or cells."""
	rows = tuple(tuple(cell_runs(cell) for cell in row.cells) for row in shape.table.rows)
	merged = any(
		cell.is_spanned or (cell.is_merge_origin and (cell.span_height > 1 or cell.span_width > 1))
		for row in shape.table.rows for cell in row.cells
	)
	unsupported_reason = "merged source table cells require span-aware emission" if merged else None
	if shape.table.first_row and rows:
		headers = rows[0]
		body_rows = rows[1:]
	else:
		headers = ()
		body_rows = rows
	return source_model.TableBlock(
		headers, body_rows, shape.left, shape.top, shape.shape_id, unsupported_reason,
	)


#============================================
def shape_inventory(
	shape: object,
	title_id: int | None,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	known_images: dict[str, str],
) -> tuple[
	list[tuple[source_model.TextRun, ...]],
	list[source_model.TextBlock],
	list[source_model.TableBlock],
	list[source_model.ImageAsset],
	list[str],
]:
	"""Extract supported semantic objects recursively from one shape."""
	titles: list[tuple[source_model.TextRun, ...]] = []
	text_blocks: list[source_model.TextBlock] = []
	tables: list[source_model.TableBlock] = []
	images: list[source_model.ImageAsset] = []
	review_reasons: list[str] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child in shape.shapes:
			child_parts = shape_inventory(child, title_id, assets_dir, djot_root, known_images)
			titles.extend(child_parts[0])
			text_blocks.extend(child_parts[1])
			tables.extend(child_parts[2])
			images.extend(child_parts[3])
			review_reasons.extend(child_parts[4])
		return titles, text_blocks, tables, images, review_reasons
	if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
		images.append(image_asset(shape, assets_dir, djot_root, known_images))
		return titles, text_blocks, tables, images, review_reasons
	if getattr(shape, "has_table", False):
		tables.append(table_block(shape))
		return titles, text_blocks, tables, images, review_reasons
	if getattr(shape, "has_text_frame", False):
		lines = tuple(
			(paragraph.level, runs)
			for paragraph in shape.text_frame.paragraphs
			if (runs := paragraph_runs(paragraph))
		)
		if not lines:
			if shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
				review_reasons.append(f"ignored non-content shape type {shape.shape_type}")
			return titles, text_blocks, tables, images, review_reasons
		if shape.shape_id == title_id:
			titles.extend(runs for _level, runs in lines)
		else:
			text_blocks.append(source_model.TextBlock(shape.left, shape.top, lines, is_subtitle_shape(shape)))
		return titles, text_blocks, tables, images, review_reasons
	if shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
		review_reasons.append(f"ignored non-content shape type {shape.shape_type}")
	return titles, text_blocks, tables, images, review_reasons


#============================================
def shape_style_evidence(shape: object) -> tuple[bool, bool]:
	"""Return direct positive fill and line evidence from one source shape."""
	properties = shape.element.find(qn("p:spPr"))
	if properties is None:
		return False, False
	fill_names = ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")
	visible_fill = any(properties.find(qn(name)) is not None for name in fill_names)
	line = properties.find(qn("a:ln"))
	visible_line = line is not None and any(line.find(qn(name)) is not None for name in fill_names)
	return visible_fill, visible_line


#============================================
def shape_placeholder_role(shape: object) -> str | None:
	"""Return a stable placeholder role independent of its source shape ID."""
	if not shape.is_placeholder:
		return None
	role = str(shape.placeholder_format.type)
	return role.split()[0].upper()


#============================================
def positioned_text_inventory(
	shape: object,
	title_id: int | None,
	z_order: tuple[int, ...] = (),
) -> list[source_model.PositionedText]:
	"""Read every text shape into raw positioned planning evidence."""
	regions: list[source_model.PositionedText] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child_index, child in enumerate(shape.shapes):
			regions.extend(positioned_text_inventory(
				child, title_id, (*z_order, child_index),
			))
		return regions
	has_positive_fill, has_positive_line = shape_style_evidence(shape)
	placeholder_role = shape_placeholder_role(shape)
	if getattr(shape, "has_table", False):
		row_count = len(shape.table.rows)
		column_count = len(shape.table.columns)
		has_header = bool(shape.table.first_row)
		merged = any(
			cell.is_spanned or (cell.is_merge_origin and (cell.span_height > 1 or cell.span_width > 1))
			for row in shape.table.rows for cell in row.cells
		)
		unsupported_reason = "merged source table cells require span-aware emission" if merged else None
		for row_index, row in enumerate(shape.table.rows):
			for column_index, cell in enumerate(row.cells):
				runs = cell_runs(cell)
				if runs:
					regions.append(source_model.PositionedText(
						((0, runs),), shape.left, shape.top, shape.width, shape.height,
						False, 0.0, False, "table",
						shape.shape_id * 10_000 + row_index * column_count + column_index,
						row_index, column_index, row_count, column_count, shape.shape_id,
						has_header, unsupported_reason, rotation_degrees=0.0,
						has_positive_fill=has_positive_fill, has_positive_line=has_positive_line,
						placeholder_role=placeholder_role, z_order=z_order,
					))
		return regions
	if not getattr(shape, "has_text_frame", False):
		return regions
	paragraphs = tuple(
		(paragraph.level, runs)
		for paragraph in shape.text_frame.paragraphs
		if (runs := paragraph_runs(paragraph))
	)
	if not paragraphs:
		return regions
	confidence = 1.0 if shape.is_placeholder else 0.0
	if shape.is_placeholder:
		source_kind = "text"
	elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
		source_kind = "text-box"
	else:
		source_kind = "auto-shape"
	regions.append(source_model.PositionedText(
		paragraphs, shape.left, shape.top, shape.width, shape.height,
		is_subtitle_shape(shape), confidence,
		shape.shape_id == title_id, source_kind, shape.shape_id,
		rotation_degrees=getattr(shape, "rotation", 0.0),
		has_positive_fill=has_positive_fill, has_positive_line=has_positive_line,
		placeholder_role=placeholder_role, z_order=z_order,
	))
	return regions


#============================================
def positioned_text_shapes(slide: object) -> tuple[source_model.PositionedText, ...]:
	"""Collect raw positioned text facts in source stack order."""
	title_id = title_shape_id(slide)
	regions: list[source_model.PositionedText] = []
	for shape_index, shape in enumerate(slide.shapes):
		regions.extend(positioned_text_inventory(shape, title_id, (shape_index,)))
	result = tuple(regions)
	return result


#============================================
def shape_has_visible_vector_content(shape: object) -> bool:
	"""Require direct shape-property fill or stroke evidence for blank vectors."""
	visible_fill, visible_line = shape_style_evidence(shape)
	return visible_fill or visible_line


#============================================
def shape_is_stroked_connector(shape: object) -> bool:
	"""Retain only OOXML-evidenced connector geometry as a separate visual role."""
	xml = shape.element.xml
	return shape_has_visible_vector_content(shape) and (
		"<p:cxnSp" in xml or 'prst="line"' in xml or "<a:headEnd" in xml or "<a:tailEnd" in xml
	)


#============================================
def connector_stroke_width(shape: object) -> float | None:
	"""Return positive raw stroke evidence for a degenerate connector."""
	properties = shape.element.find(qn("p:spPr"))
	if properties is None:
		return None
	line = properties.find(qn("a:ln"))
	width = None if line is None else line.get("w")
	if width is None or not width.isdigit():
		return None
	stroke_width = float(width)
	if stroke_width <= 0:
		return None
	return stroke_width


#============================================
def positioned_visual_inventory(
	shape: object,
	z_order: tuple[int, ...] = (),
) -> list[source_model.PositionedVisual]:
	"""Collect raw picture and non-text vector evidence for planning."""
	regions: list[source_model.PositionedVisual] = []
	if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
		for child_index, child in enumerate(shape.shapes):
			regions.extend(positioned_visual_inventory(child, (*z_order, child_index)))
		return regions
	if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
		kind = "picture"
	elif (
		shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
		and not getattr(shape, "is_placeholder", False)
		and not getattr(shape, "has_table", False)
		and (not getattr(shape, "has_text_frame", False)
			or not any(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs))
		and shape_has_visible_vector_content(shape)
	):
		kind = "connector" if shape_is_stroked_connector(shape) else "vector"
	elif not getattr(shape, "has_text_frame", False) and not getattr(shape, "has_table", False):
		kind = "connector" if shape_is_stroked_connector(shape) else "vector"
	else:
		return regions
	stroke_width = None
	if shape.width <= 0 or shape.height <= 0:
		if kind != "connector":
			return regions
		stroke_width = connector_stroke_width(shape)
		if stroke_width is None:
			return regions
	regions.append(source_model.PositionedVisual(
		f"source-{kind}-{shape.shape_id}", shape.left, shape.top, shape.width,
		shape.height, kind, shape.shape_id, stroke_width, z_order,
	))
	return regions


#============================================
def positioned_visual_shapes(slide: object) -> tuple[source_model.PositionedVisual, ...]:
	"""Return imported-deck raw visuals in source stack order."""
	visuals: list[source_model.PositionedVisual] = []
	for shape_index, shape in enumerate(slide.shapes):
		visuals.extend(positioned_visual_inventory(shape, (shape_index,)))
	return tuple(visuals)


#============================================
def slide_notes(slide: object) -> tuple[str, ...]:
	"""Read plain source notes only when the slide owns a notes part."""
	has_notes_part = any(rel.reltype.endswith("/notesSlide") for rel in slide.part.rels.values())
	if not has_notes_part:
		return ()
	notes_frame = slide.notes_slide.notes_text_frame
	if notes_frame is None:
		return ()
	return tuple(line.strip() for line in notes_frame.text.splitlines() if line.strip())


#============================================
def extract_slides(
	presentation: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	expected_hidden: set[int] | None,
) -> tuple[list[source_model.SlideData], int]:
	"""Extract raw slide facts while preserving order and authoritative visibility."""
	known_images: dict[str, str] = {}
	slides: list[source_model.SlideData] = []
	for source_index, slide in enumerate(presentation.slides, start=1):
		pptx_hidden = slide.element.get("show") == "0"
		hidden = pptx_hidden if expected_hidden is None else source_index in expected_hidden
		if expected_hidden is not None and pptx_hidden != hidden:
			raise RuntimeError(f"ODP and PPTX visibility disagree on slide {source_index}")
		if hidden:
			slides.append(source_model.SlideData(source_index, True, (), (), (), (), ()))
			continue
		titles: list[tuple[source_model.TextRun, ...]] = []
		text_blocks: list[source_model.TextBlock] = []
		tables: list[source_model.TableBlock] = []
		images: list[source_model.ImageAsset] = []
		review_reasons: list[str] = []
		title_id = title_shape_id(slide)
		for shape in slide.shapes:
			parts = shape_inventory(shape, title_id, assets_dir, djot_root, known_images)
			titles.extend(parts[0])
			text_blocks.extend(parts[1])
			tables.extend(parts[2])
			images.extend(parts[3])
			review_reasons.extend(parts[4])
		line_count = sum(len(block.lines) for block in text_blocks)
		character_count = sum(
			len(plain_text(runs)) for block in text_blocks for _level, runs in block.lines
		)
		if line_count > 10 or character_count > 1200:
			review_reasons.append("dense text requires post-conversion polish")
		text_blocks.sort(key=lambda block: (block.top, block.left))
		tables.sort(key=lambda table: (table.top, table.left))
		images.sort(key=lambda image: (image.top, image.left))
		slides.append(source_model.SlideData(
			source_index, False, tuple(titles), tuple(text_blocks), tuple(images),
			slide_notes(slide), tuple(sorted(set(review_reasons))), tuple(tables),
		))
	return slides, len(known_images)

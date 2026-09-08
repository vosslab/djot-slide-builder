"""Orchestrate bounded PPTX reading, planning, and Djot emission."""

# Standard Library
import dataclasses
import json
import os
import pathlib
import re
import shutil
import stat
import tempfile

# PIP3 modules
from pptx import Presentation

# local repo modules
import slide_lib.djot_lint as djot_lint
import slide_lib.djot_parser
import slide_lib.importers.djot_emitter as djot_emitter
import slide_lib.importers.geometry as geometry
import slide_lib.importers.pptx_reader as pptx_reader
import slide_lib.importers.slide_plan as slide_plan
import slide_lib.importers.source_model as source_model


ASSET_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]*$")


@dataclasses.dataclass(frozen=True)
class ConversionSummary:
	"""User-facing counts for one Djot conversion."""

	visible_slides: int
	editable_slides: int
	hidden_slides: int
	extracted_images: int
	review_slides: int
	output_path: pathlib.Path
	report_path: pathlib.Path


#============================================
def plan_imported_slide(
	text_regions: tuple[slide_plan.SourceTextRegion, ...],
	visual_regions: tuple[slide_plan.SourceImageRegion, ...],
	images: tuple[source_model.ImageAsset, ...],
	slide_width: int,
	slide_height: int,
) -> slide_plan.SlidePlan:
	"""Bind extracted raster assets to source geometry before semantic planning."""
	picture_sources: dict[geometry.NormalizedBounds, list[slide_plan.SourceImageRegion]] = {}
	for region in visual_regions:
		if region.source_kind == "picture":
			picture_sources.setdefault(region.bounds, []).append(region)
	extracted_images: list[slide_plan.SourceImageRegion] = []
	for image in images:
		bounds = geometry.normalized_bounds(
			image.left, image.top, image.width, image.height, slide_width, slide_height,
		)
		sources = picture_sources.get(bounds, [])
		source = sources.pop(0) if sources else None
		extracted_images.append(slide_plan.SourceImageRegion(
			image.asset_path, bounds, "picture", 0 if source is None else source.source_ordinal,
			z_order=() if source is None else source.z_order,
		))
	known_bounds = {image.bounds for image in extracted_images}
	image_regions = tuple(extracted_images) + tuple(
		image for image in visual_regions if image.bounds not in known_bounds
	)
	return slide_plan.plan_slide(text_regions, image_regions)


#============================================
def plan_slides(
	presentation: object,
	assets_dir: pathlib.Path,
	djot_root: pathlib.PurePosixPath,
	expected_hidden: set[int] | None,
) -> list[djot_emitter.PlannedSlide]:
	"""Read raw slide facts, plan their geometry, and retain both layers."""
	data_slides, _ = pptx_reader.extract_slides(
		presentation, assets_dir, djot_root, expected_hidden,
	)
	planned_slides: list[djot_emitter.PlannedSlide] = []
	visible_page_index = 0
	for source_slide, data in zip(presentation.slides, data_slides, strict=True):
		if data.hidden:
			planned_slides.append(djot_emitter.PlannedSlide(data, None))
			continue
		visible_page_index += 1
		text_regions = pptx_reader.source_text_regions(
			source_slide, presentation.slide_width, presentation.slide_height,
		)
		visual_regions = pptx_reader.source_visual_regions(
			source_slide, presentation.slide_width, presentation.slide_height,
		)
		plan = plan_imported_slide(
			text_regions,
			visual_regions,
			data.images,
			presentation.slide_width,
			presentation.slide_height,
		)
		planned_slides.append(djot_emitter.PlannedSlide(
			data, plan, text_regions, visual_regions, visible_page_index,
		))
	return planned_slides


#============================================
def validate_output_path(output_path: pathlib.Path) -> None:
	"""Protect an established Djot destination from replacement."""
	# ASVS 2.2.1 and 5.3.2: allow-list the source suffix and constrain asset placement.
	if output_path.suffix.lower() != ".djot":
		raise ValueError("output must use the .djot extension")
	if output_path.exists():
		raise FileExistsError("output Djot source already exists; import will not overwrite it")
	assets_parent = output_path.parent / "assets"
	if assets_parent.is_symlink() or (assets_parent.exists() and not assets_parent.is_dir()):
		raise ValueError("output assets parent must be a real directory")
	asset_path = assets_parent / output_path.stem
	if asset_path.exists() or asset_path.is_symlink():
		raise FileExistsError("output asset directory already exists")


#============================================
def validate_staged_djot(staging_djot: pathlib.Path, staging_assets: pathlib.Path,
		deck_stem: str) -> None:
	"""Parse, lay out, and resolve local staged assets before publication."""
	asset_target = staging_djot.parent / "assets" / deck_stem
	asset_target.parent.mkdir(exist_ok=True)
	os.symlink(staging_assets, asset_target)
	try:
		problems, _slides, _images = djot_lint.lint_source(staging_djot)
		if problems:
			problem = problems[0]
			raise ValueError(
				f"generated Djot staging validation failed: {problem.path}:{problem.line}: {problem.message}"
			)
	finally:
		asset_target.unlink(missing_ok=True)


#============================================
def staged_media_references(staging_djot: pathlib.Path, deck_stem: str) -> set[str]:
	"""Return validated direct-child media names from parsed generated Djot."""
	deck = slide_lib.djot_parser.parse_deck(staging_djot)
	expected_parent = pathlib.PurePosixPath("assets") / deck_stem
	references: set[str] = set()
	for slide in deck.slides:
		images = list(djot_lint.images_in_blocks(slide.blocks))
		for cell in slide.cells:
			images.extend(djot_lint.images_in_blocks(cell.blocks))
		for image in images:
			path = pathlib.PurePosixPath(image.source)
			if path.parent != expected_parent or len(path.parts) != 3:
				raise ValueError("generated Djot image must use its direct local asset namespace")
			if not ASSET_NAME.fullmatch(path.name) or \
				path.suffix.lower() not in pptx_reader.SUPPORTED_IMAGE_SUFFIXES:
				raise ValueError("generated Djot image has an unsafe local asset name")
			references.add(path.name)
	return references


#============================================
def prune_staged_media(staging_djot: pathlib.Path, staging_assets: pathlib.Path,
		deck_stem: str) -> int:
	"""Keep only parsed-Djot-reachable regular media in private staging assets."""
	references = staged_media_references(staging_djot, deck_stem)
	media_names: set[str] = set()
	for candidate in staging_assets.iterdir():
		metadata = candidate.lstat()
		if stat.S_ISLNK(metadata.st_mode):
			raise ValueError("staged assets must not contain symlinks")
		if stat.S_ISDIR(metadata.st_mode):
			raise ValueError("staged assets must not contain directories")
		if not stat.S_ISREG(metadata.st_mode):
			raise ValueError("staged assets must contain only regular media files")
		if not ASSET_NAME.fullmatch(candidate.name) or \
			candidate.suffix.lower() not in pptx_reader.SUPPORTED_IMAGE_SUFFIXES:
			raise ValueError("staged assets contain an unsafe media filename")
		media_names.add(candidate.name)
	if missing := references - media_names:
		raise ValueError(f"generated Djot references missing staged media: {sorted(missing)[0]}")
	for candidate in staging_assets.iterdir():
		if candidate.name not in references:
			candidate.unlink()
	return len(references)


#============================================
def publish_conversion(staging_assets: pathlib.Path, staging_djot: pathlib.Path,
		output_path: pathlib.Path) -> None:
	"""Publish both new destinations or remove only this call's partial assets."""
	assets_parent = output_path.parent / "assets"
	if assets_parent.is_symlink() or (assets_parent.exists() and not assets_parent.is_dir()):
		raise ValueError("output assets parent must be a real directory")
	assets_parent.mkdir(exist_ok=True)
	final_assets = assets_parent / output_path.stem
	if output_path.exists() or output_path.is_symlink() or \
		final_assets.exists() or final_assets.is_symlink():
		raise FileExistsError("output destination appeared during conversion; nothing was overwritten")
	if not final_assets.parent.resolve().is_relative_to(output_path.parent.resolve()):
		raise ValueError("output assets directory escapes the output parent")
	published_assets = False
	try:
		os.replace(staging_assets, final_assets)
		published_assets = True
		if output_path.exists() or output_path.is_symlink():
			raise FileExistsError("output Djot source appeared during conversion")
		os.link(staging_djot, output_path)
	except OSError:
		if published_assets:
			shutil.rmtree(final_assets)
		raise


#============================================
def convert_pptx(
	input_path: pathlib.Path,
	output_path: pathlib.Path,
	*,
	expected_slide_count: int | None = None,
	expected_hidden: set[int] | None = None,
	source_name: str | None = None,
) -> ConversionSummary:
	"""Convert one trusted PPTX into native-object-oriented Djot source."""
	input_path = input_path.resolve()
	output_path = output_path.resolve()
	# ASVS 1.5.2: validate the existing OOXML boundary before python-pptx reads it.
	pptx_reader.validate_pptx(input_path)
	validate_output_path(output_path)
	presentation = Presentation(input_path)
	if expected_slide_count is not None and len(presentation.slides) != expected_slide_count:
		raise RuntimeError("ODP and normalized PPTX slide counts disagree")
	output_path.parent.mkdir(parents=True, exist_ok=True)
	with tempfile.TemporaryDirectory(prefix=".pptx_to_djot_", dir=output_path.parent) as temporary_name:
		temporary_root = pathlib.Path(temporary_name)
		staging_assets = temporary_root / "assets"
		staging_assets.mkdir()
		djot_root = pathlib.PurePosixPath("assets") / output_path.stem
		slides = plan_slides(presentation, staging_assets, djot_root, expected_hidden)
		visible_slides = [slide for slide in slides if not slide.data.hidden]
		if not visible_slides:
			raise ValueError("presentation contains no visible slides")
		djot, records = djot_emitter.render_planned_djot(slides)
		staging_djot = temporary_root / output_path.name
		staging_djot.write_text(djot, encoding="utf-8")
		published_media_count = prune_staged_media(staging_djot, staging_assets, output_path.stem)
		report = {
			"source": source_name or input_path.name,
			"slide_count": len(slides),
			"visible_slides": len(visible_slides),
			"hidden_slides": [slide.data.source_index for slide in slides if slide.data.hidden],
			"unique_media_assets": published_media_count,
			"slides": records,
		}
		report_path = staging_assets / "import_report.json"
		report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
		validate_staged_djot(staging_djot, staging_assets, output_path.stem)
		publish_conversion(staging_assets, staging_djot, output_path)
	final_report = output_path.parent / "assets" / output_path.stem / "import_report.json"
	review_count = sum(bool(record["review_reasons"]) for record in records)
	summary = ConversionSummary(
		visible_slides=len(visible_slides),
		editable_slides=len(visible_slides),
		hidden_slides=len(slides) - len(visible_slides),
		extracted_images=published_media_count,
		review_slides=review_count,
		output_path=output_path,
		report_path=final_report,
	)
	return summary


#============================================
def run_import(input_file: pathlib.Path, output_file: pathlib.Path | None = None) -> None:
	"""Run one PPTX-to-Djot conversion."""
	output_path = input_file.with_suffix(".djot") if output_file is None else output_file
	summary = convert_pptx(input_file, output_path)
	print(
		f"Converted {summary.visible_slides} visible slides: "
		f"{summary.editable_slides} editable, {summary.review_slides} layout review, "
		f"{summary.hidden_slides} hidden, {summary.extracted_images} content images"
	)
	print(f"Djot: {summary.output_path}")
	print(f"Import report: {summary.report_path}")

#!/usr/bin/env python3
"""Exercise every Djot native layout through PPTX, ODP, and PDF export.

This E2E verifies the source-to-editable-artifact chain.  LibreOffice Impress
click-playback remains the attended presentation-fidelity gate.
"""

# Standard Library
import collections
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile

# PIP3 modules
import defusedxml.ElementTree
import PIL.Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Local Modules
import slide_lib.layout_engine
import slide_lib.layout_model
import slide_lib.layout_primitives
import slide_lib.libreoffice
import slide_lib.native_export
import slide_lib.presentation_theme


NAMESPACES = {
	"anim": "urn:oasis:names:tc:opendocument:xmlns:animation:1.0",
	"draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
	"office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
	"presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
	"style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
	"smil": "urn:oasis:names:tc:opendocument:xmlns:smil-compatible:1.0",
	"svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
	"text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


#============================================
def repo_root() -> pathlib.Path:
	"""Return the repository root containing this E2E runner."""
	result = pathlib.Path(__file__).resolve().parents[2]
	return result


#============================================
def require(condition: bool, message: str) -> None:
	"""Raise an actionable failure when an E2E contract is absent."""
	if not condition:
		raise RuntimeError(message)


#============================================
def write_image(image_path: pathlib.Path) -> None:
	"""Create the one deterministic local component image used by gallery."""
	image = PIL.Image.new("RGB", (96, 64), (36, 87, 143))
	image.save(image_path)


#============================================
def title_source(spec: slide_lib.layout_primitives.LayoutContract) -> list[str]:
	"""Return the permitted global title region for one layout specification."""
	lines: list[str] = []
	if spec.allows_title:
		lines.append(f"# {spec.name} Djot acceptance")
	if spec.allows_subtitle:
		lines.append("## Editable native objects")
	return lines


#============================================
def cell_source(spec: slide_lib.layout_primitives.LayoutContract, slot_name: str) -> list[str]:
	"""Return minimal valid source for one declared named cell."""
	if spec.name == "one-panel":
		return ["=> cascade appear", "- Native parent list item",
			"  - Native nested list item", "- Native second list item"]
	if spec.name == "multiple-choice" and slot_name == "question":
		return ["Which editable object remains visible?", "", "- Choice A", "- Choice B"]
	if spec.name == "multiple-choice" and slot_name == "answer":
		return ["Answer: Choice A"]
	if spec.name == "gallery":
		return ["![Djot gallery component one](component.png)",
			"![Djot gallery component two](component.png)"]
	return [f"Editable {spec.name} {slot_name}"]


#============================================
def slide_source(spec: slide_lib.layout_primitives.LayoutContract) -> str:
	"""Build one minimal valid Djot slide from the live layout specification."""
	lines = [f"=== layout: {spec.name}"]
	lines.extend(title_source(spec))
	for slot_name in spec.slot_names:
		lines.append(f"@{slot_name}")
		lines.extend(cell_source(spec, slot_name))
	result = "\n".join(lines)
	return result


#============================================
def write_deck(deck_path: pathlib.Path) -> tuple[str, ...]:
	"""Write one Djot slide per live layout and return their rendering order."""
	layout_names = slide_lib.layout_engine.registered_layout_names()
	source = "\n\n".join(slide_source(slide_lib.layout_engine.layout_contract(name))
		for name in layout_names) + "\n"
	deck_path.write_text(source, encoding="utf-8")
	return layout_names


#============================================
def slide_text(slide: object) -> str:
	"""Return editable text from every PPTX text frame on one slide."""
	result = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
	return result


#============================================
def inspect_multiple_choice_pptx(slide: object) -> None:
	"""Confirm question and answer are separate editable shapes in the final PPTX."""
	question_shapes = [shape for shape in slide.shapes if shape.has_text_frame and
		"Whicheditableobjectremainsvisible?" in re.sub(r"\s+", "", shape.text)]
	answer_shapes = [shape for shape in slide.shapes if shape.has_text_frame and
		"Answer:ChoiceA" in re.sub(r"\s+", "", shape.text)]
	require(len(question_shapes) == 1 and len(answer_shapes) == 1 and
		question_shapes[0] is not answer_shapes[0],
		"multiple-choice retains distinct editable question and answer shapes")


#============================================
def inspect_pptx(pptx_path: pathlib.Path, layout_names: tuple[str, ...]) -> None:
	"""Verify every Djot layout produced editable native PPTX content."""
	presentation = Presentation(pptx_path)
	require(pptx_path.stat().st_size > 0 and len(presentation.slides) == len(layout_names),
		"PPTX exists, is nonempty, and has one slide per live layout")
	for index, layout_name in enumerate(layout_names):
		if layout_name == "blank":
			continue
		text = slide_text(presentation.slides[index])
		require(bool(text.strip()), f"PPTX {layout_name} retains authored editable text")
	gallery_index = layout_names.index("gallery")
	pictures = [shape for shape in presentation.slides[gallery_index].shapes
		if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
	require(pictures, "PPTX gallery retains the Djot component image as a native picture")
	multiple_choice_index = layout_names.index("multiple-choice")
	inspect_multiple_choice_pptx(presentation.slides[multiple_choice_index])
	one_panel_index = layout_names.index("one-panel")
	list_paragraphs = {paragraph.text: paragraph
		for shape in presentation.slides[one_panel_index].shapes if shape.has_text_frame
		for paragraph in shape.text_frame.paragraphs if paragraph.text}
	parent = list_paragraphs["Native parent list item"]
	nested = list_paragraphs["Native nested list item"]
	require(parent.level == 0 and nested.level == 1 and parent._p.xml != nested._p.xml,
		"PPTX preserves distinct native list levels and theme indentation")


#============================================
def element_text(element: object) -> str:
	"""Return descendant text from one parsed ODF element."""
	result = "".join(element.itertext())
	return result


#============================================
def editable_odp_objects(page: object) -> list[object]:
	"""Return editable ODP text objects in LibreOffice's supported serializations."""
	frames = [frame for frame in page.findall(".//draw:frame", NAMESPACES)
		if frame.find(".//draw:text-box", NAMESPACES) is not None]
	custom_shapes = [shape for shape in page.findall(".//draw:custom-shape", NAMESPACES)
		if element_text(shape).strip()]
	result = frames + custom_shapes
	return result


#============================================
def classify_libreoffice_layout(layout: object) -> str:
	"""Mirror LibreOffice's ODF placeholder inference for the supported catalog."""
	object_attribute = slide_lib.presentation_theme.qname("presentation", "object")
	x_attribute = slide_lib.presentation_theme.qname("svg", "x")
	placeholders = layout.findall("./presentation:placeholder", NAMESPACES)
	objects = tuple(item.attrib[object_attribute] for item in placeholders)
	x_values = tuple(float(item.attrib[x_attribute].removesuffix("cm")) for item in placeholders)
	if not objects:
		return "AUTOLAYOUT_NONE"
	if len(objects) == 1:
		return "AUTOLAYOUT_TITLE_ONLY" if objects[0] == "title" else "AUTOLAYOUT_ONLY_TEXT"
	if len(objects) == 2:
		if objects[1] == "subtitle":
			return "AUTOLAYOUT_TITLE"
		if objects[1] == "outline":
			return "AUTOLAYOUT_TITLE_CONTENT"
		if objects[1] == "vertical_outline":
			if objects[0] == "vertical_title":
				return "AUTOLAYOUT_VTITLE_VCONTENT"
			return "AUTOLAYOUT_TITLE_VCONTENT"
		return "AUTOLAYOUT_NONE"
	if len(objects) == 3:
		if objects[1] == "outline" and objects[2] == "outline":
			return "AUTOLAYOUT_TITLE_2CONTENT"
		if objects[1] == "graphic" and objects[2] == "vertical_outline":
			return "AUTOLAYOUT_TITLE_2VTEXT"
		if objects[1] == "vertical_outline":
			return "AUTOLAYOUT_VTITLE_VCONTENT_OVER_VCONTENT"
		if objects[1] not in ("outline", "chart", "graphic") and x_values[1] >= x_values[2]:
			return "AUTOLAYOUT_TITLE_CONTENT_OVER_CONTENT"
		return "AUTOLAYOUT_NONE"
	if len(objects) == 4:
		if objects[1] != "object":
			return "AUTOLAYOUT_TITLE_CONTENT_2CONTENT"
		if x_values[1] < x_values[2]:
			return "AUTOLAYOUT_TITLE_2CONTENT_OVER_CONTENT"
		return "AUTOLAYOUT_TITLE_2CONTENT_CONTENT"
	if len(objects) == 5 and objects[1] == "object":
		return "AUTOLAYOUT_TITLE_4CONTENT"
	if len(objects) == 7:
		return "AUTOLAYOUT_TITLE_6CONTENT"
	return "AUTOLAYOUT_NONE"


#============================================
def reveal_target_texts(page: object) -> tuple[str, ...]:
	"""Interpret native click order as the editable text each step reveals."""
	xml_id = "{http://www.w3.org/XML/1998/namespace}id"
	target_attribute = f"{{{NAMESPACES['smil']}}}targetElement"
	elements = {element.attrib[xml_id]: element for element in page.iter()
		if xml_id in element.attrib}
	targets = [target for element in page.findall(".//anim:set", NAMESPACES) +
		page.findall(".//anim:transitionFilter", NAMESPACES)
		if (target := element.attrib.get(target_attribute)) is not None]
	return tuple(element_text(elements[target]) for target in targets)


#============================================
def inspect_odp(odp_path: pathlib.Path, layout_names: tuple[str, ...],
		plan: slide_lib.layout_model.LayoutDeck, require_distinct_layouts: bool = True,
		after_libreoffice: bool = False) -> None:
	"""Verify LibreOffice preserved separate editable ODP text and image objects."""
	with zipfile.ZipFile(odp_path) as archive:
		content_root = defusedxml.ElementTree.fromstring(archive.read("content.xml"))
		styles_root = defusedxml.ElementTree.fromstring(archive.read("styles.xml"))
	pages = content_root.findall("./office:body/office:presentation/draw:page", NAMESPACES)
	require(odp_path.stat().st_size > 0 and len(pages) == len(layout_names),
		"ODP exists, is nonempty, and has one page per live layout")
	theme = slide_lib.presentation_theme.default_theme()
	master_attribute = slide_lib.presentation_theme.qname("draw", "master-page-name")
	masters = styles_root.findall(".//style:master-page", NAMESPACES)
	require(all(page.attrib[master_attribute] == theme.master_name for page in pages) and
		len(masters) == 1,
		"ODP pages use the sole authoritative template master")
	layout_attribute = slide_lib.presentation_theme.qname(
		"presentation", "presentation-page-layout-name")
	name_attribute = slide_lib.presentation_theme.qname("style", "name")
	defined_layouts = {item.attrib[name_attribute] for item in
		styles_root.findall(".//style:presentation-page-layout", NAMESPACES)}
	referenced_values = [page.attrib.get(layout_attribute) for page in pages]
	missing_layouts = [layout_names[index] for index, value in enumerate(referenced_values)
		if value is None and layout_names[index] != "blank"]
	referenced_layouts = {value for value in referenced_values if value is not None}
	require(not missing_layouts and referenced_layouts <= defined_layouts and
		(not require_distinct_layouts or len(referenced_layouts) == len(layout_names)),
		"each nonblank Djot layout resolves to one native LibreOffice page-layout " +
		f"definition; missing {missing_layouts}")
	layouts = {item.attrib[name_attribute]: item for item in
		styles_root.findall(".//style:presentation-page-layout", NAMESPACES)}
	standard_identities = tuple(contract.libreoffice_autolayout for name in layout_names
		if (contract := slide_lib.layout_engine.layout_contract(name)).libreoffice_autolayout
		is not None)
	require(len(standard_identities) == len(set(standard_identities)),
		"standard Djot layouts declare distinct LibreOffice AutoLayout identities")
	classification_mismatches: list[str] = []
	for index, layout_name in enumerate(layout_names):
		contract = slide_lib.layout_engine.layout_contract(layout_name)
		if contract.libreoffice_autolayout is None or (after_libreoffice and layout_name == "blank"):
			continue
		layout_reference = referenced_values[index]
		actual = "missing" if layout_reference is None else classify_libreoffice_layout(
			layouts[layout_reference])
		expected = contract.libreoffice_autolayout.value
		if actual != expected:
			classification_mismatches.append(f"{layout_name}: expected {expected}, found {actual}")
	require(not classification_mismatches,
		"each standard Djot layout selects its distinct LibreOffice AutoLayout; " +
		"; ".join(classification_mismatches))
	class_attribute = slide_lib.presentation_theme.qname("presentation", "class")
	page_occupancies = tuple(collections.Counter(item.attrib[class_attribute] for item in list(page)
		if class_attribute in item.attrib) for page in pages)
	expected_occupancies = tuple(collections.Counter(member.role.value for member in
		slide.layout.topology.members) for slide in plan.slides)
	differences = tuple(f"{layout_names[index]}: expected {expected}, found {page_occupancies[index]}"
		for index, expected in enumerate(expected_occupancies)
		if (not after_libreoffice or layout_names[index] != "multiple-choice") and
		page_occupancies[index] != expected)
	require(not differences,
		"each ODP page occupies its compiled editable presentation members exactly once; " +
		"; ".join(differences))
	xml_id = "{http://www.w3.org/XML/1998/namespace}id"
	target_attribute = f"{{{NAMESPACES['smil']}}}targetElement"
	ids = {element.attrib[xml_id] for element in content_root.iter() if xml_id in element.attrib}
	timing_nodes = content_root.findall(".//anim:set", NAMESPACES) + \
		content_root.findall(".//anim:transitionFilter", NAMESPACES)
	targets = [target for element in timing_nodes
		if (target := element.attrib.get(target_attribute)) is not None]
	require(len(targets) >= 3 and set(targets) <= ids,
		"native reveal steps retain resolvable ODF target identities")
	direct_bands = [shape for page in pages
		for shape in page.findall("./draw:custom-shape", NAMESPACES)
		if shape.attrib.get(slide_lib.presentation_theme.qname("svg", "x")) == "0cm" and
		shape.attrib.get(slide_lib.presentation_theme.qname("svg", "y")) == "0cm" and
		not element_text(shape).strip()]
	require(not direct_bands, "ODP top band comes from the master rather than repeated slide shapes")
	one_panel_index = layout_names.index("one-panel")
	parent_outline = pages[one_panel_index].find(".//text:list", NAMESPACES)
	nested_outline = None if parent_outline is None else parent_outline.find(
		"./text:list-item/text:list", NAMESPACES)
	require(parent_outline is not None and nested_outline is not None and
		"Native nested list item" in element_text(nested_outline),
		"ODP preserves parent and nested items as native level-specific lists")
	one_panel_reveals = reveal_target_texts(pages[one_panel_index])
	require(len(one_panel_reveals) == 2 and
		"Native parent list item" in one_panel_reveals[0] and
		"Native second list item" in one_panel_reveals[1],
		"ODP cascade click order retains the two top-level source items")
	gallery_index = layout_names.index("gallery")
	gallery_images = pages[gallery_index].findall(".//draw:image", NAMESPACES)
	require(gallery_images, "ODP gallery retains the Djot component image as an editable draw image")
	multiple_choice_index = layout_names.index("multiple-choice")
	multiple_choice_objects = editable_odp_objects(pages[multiple_choice_index])
	question_objects = [item for item in multiple_choice_objects
		if "Which editable object remains visible?" in element_text(item)]
	answer_objects = [item for item in multiple_choice_objects
		if "Answer: Choice A" in element_text(item)]
	require(len(question_objects) == 1 and len(answer_objects) == 1 and
		question_objects[0] is not answer_objects[0],
		"ODP retains multiple-choice question and answer in distinct editable text objects")
	answer_reveals = reveal_target_texts(pages[multiple_choice_index])
	answer_timing = pages[multiple_choice_index].find(".//anim:set", NAMESPACES)
	answer_target = None if answer_timing is None else answer_timing.attrib.get(target_attribute)
	answer_element = next((item for item in pages[multiple_choice_index].iter()
		if item.attrib.get(xml_id) == answer_target), None)
	require(len(answer_reveals) == 1 and "Answer: Choice A" in answer_reveals[0] and
		answer_element is not None and answer_element.tag ==
		slide_lib.presentation_theme.qname("draw", "frame"),
		"ODP answer popup retains its authored text and native frame reveal target")


#============================================
def inspect_pdf(pdf_path: pathlib.Path, layout_names: tuple[str, ...]) -> None:
	"""Confirm the ODP-derived PDF is nonempty and has one page per layout."""
	result = subprocess.run(["pdfinfo", str(pdf_path)], check=True, capture_output=True, text=True)
	page_match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, re.MULTILINE)
	require(pdf_path.stat().st_size > 0 and page_match is not None and
		int(page_match.group(1)) == len(layout_names),
		"ODP-derived PDF is nonempty and has one page per live layout")


#============================================
def run() -> None:
	"""Run source Djot through the public PPTX, ODP, and PDF export chain."""
	root = repo_root()
	output_root = root / "output"
	output_root.mkdir(exist_ok=True)
	stem = f"djot_native_layouts_{uuid.uuid4().hex}"
	workspace = pathlib.Path(tempfile.mkdtemp(prefix=f"{stem}_", dir=output_root))
	deck_path = workspace / f"{stem}.djot"
	pptx_path = output_root / "pptx" / f"{stem}.pptx"
	odp_path = output_root / "odp" / f"{stem}.odp"
	pdf_path = output_root / "pdf" / f"{stem}.pdf"
	completed = False
	try:
		write_image(workspace / "component.png")
		layout_names = write_deck(deck_path)
		plan, _theme = slide_lib.native_export.compile_deck(
			slide_lib.native_export.parse_deck(deck_path))
		command = [sys.executable, "deck_tools.py", "build", str(deck_path), "--format", "all"]
		subprocess.run(command, cwd=root, check=True)
		inspect_pptx(pptx_path, layout_names)
		inspect_odp(odp_path, layout_names, plan)
		inspect_pdf(pdf_path, layout_names)
		fodp_root = workspace / "fodp"
		round_trip_root = workspace / "round_trip"
		fodp_root.mkdir()
		round_trip_root.mkdir()
		fodp_path = slide_lib.libreoffice.convert_file(odp_path, fodp_root, "fodp")
		round_trip_path = slide_lib.libreoffice.convert_file(fodp_path, round_trip_root, "odp")
		inspect_odp(round_trip_path, layout_names, plan, require_distinct_layouts=False,
			after_libreoffice=True)
		print("PASS: Djot source retains editable PPTX and ODP semantics through LibreOffice and PDF")
		completed = True
	finally:
		if completed:
			for artifact_path in (pptx_path, odp_path, pdf_path):
				artifact_path.unlink(missing_ok=True)
			shutil.rmtree(workspace, ignore_errors=True)
		else:
			print(f"FAILED: diagnostic artifacts retained under {workspace} and {output_root}",
				file=sys.stderr)


if __name__ == "__main__":
	run()

"""Behavioral tests for PPTX-to-extended-Djot conversion."""

# Standard Library
import pathlib
import types

# PIP3 modules
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

# Local modules
import slide_lib.djot_parser
import slide_lib.importers.djot_emitter as djot_emitter
import slide_lib.importers.geometry as geometry
import slide_lib.importers.pptx_reader as pptx_reader
import slide_lib.importers.slide_plan as slide_plan
import slide_lib.importers.pptx_to_djot as pptx_to_djot
import slide_lib.importers.source_model as source_model
import slide_lib.layout_engine
import slide_lib.layout_validation
import slide_lib.native_model
#============================================
def write_png(output_path: pathlib.Path) -> pathlib.Path:
	"""Write one bounded source image."""
	image = Image.new("RGB", (40, 30), (20, 90, 160))
	image.save(output_path)
	return output_path
#============================================
def text_region(
	left: float,
	top: float,
	right: float,
	bottom: float,
	text: str,
	*,
	title_identity: bool = False,
	placeholder_confidence: float = 0.0,
	source_kind: str = "text",
	paragraph_count: int = 1,
	source_ordinal: int = 0,
	z_order: tuple[int, ...] = (),
) -> slide_plan.SourceTextRegion:
	"""Build one concise synthetic source text region."""
	return slide_plan.SourceTextRegion(
		tuple((0, (source_model.TextRun(text),)) for _index in range(paragraph_count)),
		geometry.NormalizedBounds(left, top, right, bottom),
		False,
		1.0 if title_identity else placeholder_confidence,
		title_identity,
		source_kind,
		source_ordinal,
		z_order=z_order,
	)
#============================================
def test_empty_notes_part_is_treated_as_no_notes() -> None:
	"""An empty LibreOffice notes relationship never aborts a Djot import."""
	notes_relation = types.SimpleNamespace(reltype="http://example.test/notesSlide")
	slide = types.SimpleNamespace(
		part=types.SimpleNamespace(rels={"notes": notes_relation}),
		notes_slide=types.SimpleNamespace(notes_text_frame=None),
	)
	assert pptx_reader.slide_notes(slide) == ()


#============================================
def test_paragraph_runs_preserve_characters_across_style_boundaries() -> None:
	"""Reader runs never invent whitespace at a source formatting boundary."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	paragraph = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.paragraphs[0]
	first = paragraph.add_run()
	first.text = "bio"
	first.font.bold = True
	second = paragraph.add_run()
	second.text = "chemistry"

	runs = pptx_reader.paragraph_runs(paragraph)

	assert tuple(run.text for run in runs) == ("bio", "chemistry")
#============================================
def test_blank_auto_shape_inventory_requires_visible_vector_evidence() -> None:
	"""Only blank auto-shapes with OOXML fill or stroke enter visual inventory."""
	def shape(*, fill: bool = False, stroke: bool = False, xml: str = "<p:sp>") -> object:
		line = types.SimpleNamespace(find=lambda name: object() if stroke and name.endswith("solidFill") else None)
		properties = types.SimpleNamespace(find=lambda name: line if name.endswith("ln") else (
			object() if fill and name.endswith("Fill") else None))
		return types.SimpleNamespace(
			shape_type=MSO_SHAPE_TYPE.AUTO_SHAPE, is_placeholder=False,
			has_text_frame=True, has_table=False,
			text_frame=types.SimpleNamespace(paragraphs=()),
			element=types.SimpleNamespace(xml=xml, find=lambda name: properties if name.endswith("spPr") else None), shape_id=7,
			left=0, top=0, width=20, height=20,
		)
	visible = shape(fill=True)
	invisible = shape(xml="<p:txBody><a:rPr><a:solidFill/></a:rPr></p:txBody>")
	assert pptx_reader.positioned_visual_inventory(visible)
	assert not pptx_reader.positioned_visual_inventory(invisible)
#============================================
def test_vector_media_is_validated_and_emf_type_is_normalized() -> None:
	"""Older vector blobs retain their actual type only after header validation."""
	emf_blob = b"\x01\x00\x00\x00" + b"\x00" * 36 + b" EMF" + b"\x00" * 8

	pptx_reader.validate_image_blob(emf_blob, ".emf")
	assert pptx_reader.image_suffix(emf_blob, ".wmf") == ".emf"
	with pytest.raises(ValueError, match="WMF image header"):
		pptx_reader.validate_image_blob(b"not a metafile", ".wmf")
#============================================
def test_dna_text_preserves_libreoffice_prime_and_split_run_notation() -> None:
	"""Observed source variants survive normalization and formatting-only run splits."""
	assert djot_emitter.djot_text("5'-AGTACT-'3") == "5&prime;-`AGTACT`-3&prime;"
	assert djot_emitter.render_runs((
		source_model.TextRun("5'-G|"), source_model.TextRun("AATTC-3'"),
	)) == "5&prime;-`G|AATTC`-3&prime;"
#============================================
def test_converter_refuses_non_djot_or_existing_output(tmp_path: pathlib.Path) -> None:
	"""The bounded importer never overwrites another source format."""
	with pytest.raises(ValueError, match=".djot"):
		pptx_to_djot.validate_output_path(tmp_path / "lecture.md")
	output_path = tmp_path / "lecture.djot"
	output_path.write_text("existing\n", encoding="utf-8")
	with pytest.raises(FileExistsError, match="will not overwrite"):
		pptx_to_djot.validate_output_path(output_path)
#============================================
def test_geometry_beats_misleading_title_placeholder_identity() -> None:
	"""A shallow wide heading wins over an identity-tagged body placeholder."""
	shallow = text_region(
		0.08, 0.06, 0.88, 0.16, "Actual title", placeholder_confidence=1.0,
	)
	misleading = text_region(0.14, 0.44, 0.78, 0.70, "Body placeholder", title_identity=True)
	decision = slide_plan.select_title((misleading, shallow))
	assert decision.region is shallow
	assert decision.reason == "shallow-wide geometry"
#============================================
def test_unmarked_diagram_label_and_tall_identity_stay_out_of_title_lane() -> None:
	"""Diagram labels, body placeholders, and side titles stay out of the H1 lane."""
	label = text_region(0.08, 0.04, 0.92, 0.18, "Diagram pathway", source_kind="text-box")
	identity = text_region(0.05, 0.25, 0.95, 0.80, "Body", title_identity=True)
	side_title = text_region(0.05, 0.04, 0.40, 0.80, "Side title", title_identity=True)
	decision = slide_plan.select_title((label, identity))
	solitary_decision = slide_plan.select_solitary_title((side_title,), ())

	assert decision.region is None
	assert solitary_decision is None
#============================================
def test_solitary_tall_title_identity_becomes_one_global_title() -> None:
	"""One sole text placeholder remains a title despite a tall template box."""
	region = slide_plan.SourceTextRegion(
		((0, (source_model.TextRun("Course"),)),
			(0, (source_model.TextRun("Introduction"),))),
		geometry.NormalizedBounds(0.05, 0.04, 0.95, 0.88),
		True,
		1.0,
		True,
	)

	plan = slide_plan.plan_slide((region,), ())

	assert plan.title.region is region
#============================================
def test_geometry_partitions_clear_left_and_right_text_columns() -> None:
	"""Separated text components become named slots in reading order."""
	left = text_region(0.06, 0.28, 0.40, 0.72, "Left explanation", paragraph_count=5)
	right = text_region(0.61, 0.28, 0.94, 0.72, "Right explanation", paragraph_count=5)
	image = slide_plan.SourceImageRegion(
		"assets/right.png",
		geometry.NormalizedBounds(0.70, 0.30, 0.90, 0.50),
	)

	plan = slide_plan.plan_slide((left, right), (image,))

	assert tuple(slot.name for slot in plan.slots) == ("left", "right")
#============================================
def test_single_auto_shape_caption_keeps_exterior_caption_native() -> None:
	"""One interior auto-shape diagram label leaves an exterior caption editable."""
	image = slide_plan.SourceImageRegion(
		"assets/diagram.png", geometry.NormalizedBounds(0.05, 0.10, 0.85, 0.80),
	)
	interior = text_region(0.35, 0.40, 0.45, 0.48, "Interior", source_kind="auto-shape")
	exterior = text_region(0.88, 0.40, 0.98, 0.48, "Exterior", source_kind="auto-shape")
	plan = slide_plan.plan_slide((interior, exterior), (image,))

	assert plan.content_region is not None
	assert plan.slots[0].text_regions == (exterior,)
#============================================
def test_complete_placeholder_list_and_callout_emit_multiple_choice() -> None:
	"""A complete question-and-answer structure emits answer semantics."""
	question = slide_plan.SourceTextRegion(
		((0, (source_model.TextRun("Prompt"),)),
			(1, (source_model.TextRun("Choice one"),)),
			(1, (source_model.TextRun("Choice two"),))),
		geometry.NormalizedBounds(0.08, 0.20, 0.82, 0.68), False, 1.0,
		source_ordinal=4,
	)
	answer = slide_plan.SourceTextRegion(
		((0, (source_model.TextRun("Answer"),)),),
		geometry.NormalizedBounds(0.60, 0.63, 0.80, 0.67),
		False, 0.0, source_ordinal=8,
	)
	plan = slide_plan.plan_slide((question, answer), ())
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(4, False, (), (), (), (), ()), plan, visible_page_index=3,
	)

	lines, layout, _reasons = djot_emitter.render_planned_slide(planned, False)

	assert layout == "multiple-choice"
	assert "@question" in lines and "@answer" in lines


#============================================
def test_regular_text_grid_stays_reviewable_without_table_metadata() -> None:
	"""Text centers alone never claim a faithful native table."""
	regions = tuple(
		text_region(
			0.05 + column * 0.25,
			0.20 + row * 0.18,
			0.20 + column * 0.25,
			0.26 + row * 0.18,
			"cell",
			source_kind="auto-shape",
		)
		for row in range(3)
		for column in range(3)
	)

	plan = slide_plan.plan_slide(regions, ())

	assert plan.tables == ()
	assert plan.review_reason is not None
#============================================
def test_actual_pptx_table_preserves_blank_cells(tmp_path: pathlib.Path) -> None:
	"""A true PPTX table retains dimensions and an intentional blank cell."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
	table.first_row = False
	table.cell(0, 0).text = "Header"
	table.cell(1, 1).text = "Value"
	slide.shapes.add_picture(str(write_png(tmp_path / "backing.png")), Inches(0.1), Inches(0.1))
	path = tmp_path / "table.pptx"
	presentation.save(path)
	loaded = Presentation(path)
	loaded_slide = loaded.slides[0]

	plan = pptx_to_djot.plan_imported_slide(
		slide_plan.text_regions(
			pptx_reader.positioned_text_shapes(loaded_slide),
			loaded.slide_width, loaded.slide_height,
		),
		slide_plan.visual_regions(
			pptx_reader.positioned_visual_shapes(loaded_slide),
			loaded.slide_width, loaded.slide_height,
		),
		(), loaded.slide_width, loaded.slide_height,
	)

	assert plan.tables[0].rows[0][1].text_regions == ()
	assert plan.content_region is None


#============================================
def test_imported_pptx_table_parses_back_as_a_native_table(tmp_path: pathlib.Path) -> None:
	"""Raw table facts survive one Djot emission and parsing round trip."""
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
	table.first_row = True
	table.cell(0, 0).text = "Condition"
	table.cell(0, 1).text = "Source"
	table.cell(1, 0).text = "A | B"
	table.cell(1, 1).text = "5'-AGTACT-3'"
	table.cell(0, 1).text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.test/source"
	input_path = tmp_path / "table.pptx"
	output_path = tmp_path / "table.djot"
	presentation.save(input_path)

	pptx_to_djot.convert_pptx(input_path, output_path)
	deck = slide_lib.djot_parser.parse_deck(output_path)
	parsed = next(block for block in deck.slides[0].cells[0].blocks
		if isinstance(block, slide_lib.native_model.Table))
	source_link = parsed.headers[1][0]

	assert isinstance(source_link, slide_lib.native_model.Link) and \
		source_link.url == "https://example.test/source"
	assert tuple(slide_lib.djot_parser.visible_text(cell) for cell in (
		parsed.headers[0], parsed.rows[0][0], parsed.rows[0][1],
	)) == ("Condition", "A | B", "5\u2032-AGTACT-3\u2032")
#============================================
def test_side_by_side_text_and_image_use_live_two_panels() -> None:
	"""A live side-by-side topology keeps direct text and image components separate."""
	prose = text_region(.08, .32, .45, .76, "Editable prose")
	image = slide_plan.SourceImageRegion("assets/photo.png", geometry.NormalizedBounds(.55, .32, .92, .76))

	plan = slide_plan.plan_slide((prose,), (image,))

	assert tuple(slot.name for slot in plan.slots) == ("left", "right")


#============================================
def test_genuine_vertical_lane_uses_one_ordered_flow() -> None:
	"""A source-ordered bullet, picture, and caption share one vertical lane."""
	bullet = text_region(.10, .35, .90, .42, "Editable bullet")
	caption = text_region(.10, .75, .90, .80, "Editable caption")
	image = slide_plan.SourceImageRegion("assets/photo.png", geometry.NormalizedBounds(.22, .40, .78, .70))

	plan = slide_plan.plan_slide((bullet, caption), (image,))

	assert plan.slots[0].flows_in_source_order
	assert plan.slots[0].text_regions == (bullet, caption)


#============================================
def test_planned_coupled_content_normalizes_to_native_flow() -> None:
	"""The emitter keeps coupled labels as native content in a standard cell."""
	title = text_region(0.05, 0.04, 0.90, 0.14, "Editable title", title_identity=True)
	left = text_region(0.05, 0.30, 0.35, 0.70, "Native left")
	annotation = text_region(0.65, 0.42, 0.72, 0.48, "Diagram label")
	image = slide_plan.SourceImageRegion(
		"assets/diagram.png", geometry.NormalizedBounds(0.45, 0.25, 0.92, 0.82),
	)
	content = slide_plan.ContentRegionPlan(
		"content-region-1", geometry.NormalizedBounds(0.45, 0.25, 0.92, 0.82),
		(annotation,), (image,),
	)
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(title, "test"),
		(slide_plan.SlotPlan("left", (left,)), slide_plan.SlotPlan("right", ())),
		content,
	)
	data = source_model.SlideData(1, False, (), (), (), (), ())
	planned = djot_emitter.PlannedSlide(data, plan, (title, left, annotation), (image,), 1)

	lines, layout, reasons = djot_emitter.render_planned_slide(planned, False)

	assert layout == "two-panels"
	assert "- Diagram label" in lines
	assert any("normalized" in reason for reason in reasons)


#============================================
def test_planned_content_without_native_members_emits_review_placeholder() -> None:
	"""An unsupported legacy relationship remains visible without a raster substitute."""
	content = slide_plan.ContentRegionPlan(
		"content-region-1", geometry.NormalizedBounds(0.20, 0.20, 0.80, 0.80), (), (),
	)
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(None, "test"),
		(slide_plan.SlotPlan("body", ()),), content,
	)
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	lines, layout, reasons = djot_emitter.render_planned_slide(planned, False)
	assert layout == "one-panel" and any("Native reconstruction needed" in line for line in lines)
	assert any("normalized" in reason for reason in reasons)


#============================================
def test_separable_images_emit_once_as_atomic_components() -> None:
	"""Each source shape owns one image component, even for a shared asset."""
	first = slide_plan.SourceImageRegion(
		"assets/shared.png", geometry.NormalizedBounds(0.05, 0.25, 0.40, 0.70),
		source_ordinal=11,
	)
	second = slide_plan.SourceImageRegion(
		"assets/shared.png", geometry.NormalizedBounds(0.55, 0.25, 0.90, 0.70),
		source_ordinal=12,
	)
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(None, "test"),
		(slide_plan.SlotPlan("body", (), (first, second)),),
	)
	data = source_model.SlideData(
		1, False, (), (), (
			source_model.ImageAsset(1, 1, 1, 1, "assets/shared.png", "Shared"),
		), (), (),
	)
	planned = djot_emitter.PlannedSlide(data, plan, visible_page_index=1)

	components, _reasons = djot_emitter.emit_components(planned)
	assert tuple(component.source_image_ids for component in components) == (
		((11, "assets/shared.png"),), ((12, "assets/shared.png"),),
	)
	assert djot_emitter.render_planned_slide(planned, False)[1] == "two-panels"


#============================================
def test_table_projection_keeps_blank_cells_editable() -> None:
	"""Supported rectangular source tables become canonical editable pipe tables."""
	filled = text_region(0.10, 0.30, 0.30, 0.40, "Header", source_kind="table")
	blank = slide_plan.TableCellPlan(0, 1, ())
	table = slide_plan.TablePlan(
		geometry.NormalizedBounds(0.10, 0.30, 0.70, 0.70), (filled,),
		((slide_plan.TableCellPlan(0, 0, (filled,)), blank),), 1, 2,
	)
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(None, "test"),
		(slide_plan.SlotPlan("body", ()),), tables=(table,),
	)
	raw_table = source_model.TableBlock(
		((source_model.TextRun("Header"),), (source_model.TextRun(""),)), (), 0, 0, 0,
	)
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(1, False, (), (), (), (), (), (raw_table,)),
		plan, visible_page_index=1,
	)

	lines, _layout, reasons = djot_emitter.render_planned_slide(planned, False)

	assert "| Header |  |" in lines
	assert reasons == []


#============================================
def test_table_fallback_keeps_table_and_prose_in_separate_native_cells(
	tmp_path: pathlib.Path,
) -> None:
	"""Overlapping table and prose normalize to valid editable regions."""
	prose = text_region(0.10, 0.30, 0.70, 0.70, "Editable prose", source_ordinal=4,
		z_order=(0,))
	filled = text_region(0.10, 0.30, 0.70, 0.70, "Header", source_kind="table",
		source_ordinal=50_000, z_order=(1,))
	table = slide_plan.TablePlan(
		filled.bounds, (filled,), ((slide_plan.TableCellPlan(0, 0, (filled,)),),),
		1, 1, table_id=5,
	)
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(None, "test"),
		(slide_plan.SlotPlan("body", (prose,)),), tables=(table,),
	)
	raw_table = source_model.TableBlock(((source_model.TextRun("Header"),),), (), 0, 0, 5)
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(1, False, (), (), (), (), (), (raw_table,)),
		plan, visible_page_index=1,
	)

	lines, layout, _reasons = djot_emitter.render_planned_slide(planned, False)
	path = tmp_path / "normalized.djot"
	path.write_text("\n".join(lines) + "\n", encoding="utf-8")
	parsed = slide_lib.djot_parser.parse_deck(path).slides[0]
	slide_lib.layout_validation.validate_layout_source(
		parsed, slide_lib.layout_engine.layout_contract(layout),
	)
	cell_blocks = tuple(block for cell in parsed.cells for block in cell.blocks)

	assert layout == "two-panels"
	assert sum(isinstance(block, slide_lib.native_model.Table) for block in cell_blocks) == 1


#============================================
def test_one_panel_fallback_uses_source_stack_order() -> None:
	"""Geometry fallback keeps authored stack order instead of visual coordinates."""
	first = text_region(.40, .40, .90, .90, "First", source_ordinal=10, z_order=(0,))
	second = text_region(.10, .10, .60, .60, "Second", source_ordinal=11, z_order=(1,))
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(None, "test"),
		(slide_plan.SlotPlan("body", (first, second)),),
	)
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	lines, layout, _reasons = djot_emitter.render_planned_slide(planned, False)

	assert layout == "one-panel"
	assert lines.index("- First") < lines.index("- Second")


#============================================
def test_import_report_retains_normalized_text_and_links() -> None:
	"""Review records preserve source facts omitted from visible reconstruction."""
	label = slide_plan.SourceTextRegion(
		((0, (source_model.TextRun("A & B", "https://example.test/label"),)),),
		geometry.NormalizedBounds(.20, .20, .30, .30), False, 0.0,
		source_kind="auto-shape", source_ordinal=7, z_order=(2,),
	)
	content = slide_plan.ContentRegionPlan(
		"content-region-1", label.bounds, (label,), (),
	)
	plan = slide_plan.SlidePlan(slide_plan.TitleDecision(None, "test"), (), content)
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(1, False, (), (), (), (), ()), plan, visible_page_index=1,
	)

	_record_source, records = djot_emitter.render_planned_djot([planned])
	runs = records[0]["content_region"]["text_regions"][0]["paragraphs"][0]["runs"]

	assert runs[0]["text"] == "A & B"
	assert runs[0]["link"] == "https://example.test/label"


#============================================
def test_import_report_preserves_presenter_note_text() -> None:
	"""A migration report retains notes that Djot cannot yet author."""
	data = source_model.SlideData(
		1, False, (), (), (), ("Explain the diagram after revealing the answer.",), (),
	)
	planned = djot_emitter.PlannedSlide(
		data, slide_plan.SlidePlan(slide_plan.TitleDecision(None, "test"), ()),
		visible_page_index=1,
	)

	_source, records = djot_emitter.render_planned_djot([planned])

	assert records[0]["presenter_notes"] == [
		"Explain the diagram after revealing the answer.",
	]


#============================================
def test_merged_table_requires_review_instead_of_flattening() -> None:
	"""A span unsupported by the native table contract fails before publication."""
	table = slide_plan.TablePlan(
		geometry.NormalizedBounds(0.10, 0.30, 0.70, 0.70), (), (), 1, 1,
		unsupported_reason="merged cells are not representable",
	)
	plan = slide_plan.SlidePlan(
		slide_plan.TitleDecision(None, "test"),
		(slide_plan.SlotPlan("body", ()),), tables=(table,),
	)
	raw_table = source_model.TableBlock((), (), 0, 0, 0, "merged cells are not representable")
	planned = djot_emitter.PlannedSlide(
		source_model.SlideData(1, False, (), (), (), (), (), (raw_table,)),
		plan, visible_page_index=1,
	)

	with pytest.raises(ValueError, match="source table requires review"):
		djot_emitter.render_planned_slide(planned, False)


#============================================
def test_planner_body_and_coupled_component_receive_native_cells() -> None:
	"""A normal body plus legacy diagram content emits as native standard cells."""
	title = text_region(0.05, 0.04, 0.90, 0.14, "Title", title_identity=True)
	prose = text_region(0.05, 0.30, 0.35, 0.75, "Editable prose")
	first = text_region(0.55, 0.35, 0.62, 0.42, "Label one", source_kind="auto-shape")
	second = text_region(0.75, 0.62, 0.82, 0.69, "Label two", source_kind="auto-shape")
	image = slide_plan.SourceImageRegion(
		"assets/diagram.png", geometry.NormalizedBounds(0.45, 0.25, 0.92, 0.82),
	)
	plan = slide_plan.plan_slide((title, prose, first, second), (image,))
	data = source_model.SlideData(1, False, (), (), (
		source_model.ImageAsset(0, 0, 1, 1, "assets/diagram.png", "Diagram"),
	), (), ())
	planned = djot_emitter.PlannedSlide(data, plan, visible_page_index=1)

	lines, layout, reasons = djot_emitter.render_planned_slide(planned, False)

	assert layout == "two-panels" and "- Editable prose" in lines
	assert "- Label one" in lines and "![Diagram](assets/diagram.png)" in lines
	assert any("normalized" in reason for reason in reasons)


#============================================
def test_publish_rollback_removes_only_just_published_assets(
	tmp_path: pathlib.Path,
	monkeypatch: pytest.MonkeyPatch,
) -> None:
	"""A failed second publication step leaves neither new output destination behind."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	(staging_assets / "image.png").write_bytes(b"image")
	staging_djot = tmp_path / "staging.djot"
	staging_djot.write_text("source\n", encoding="utf-8")
	output_path = tmp_path / "deck.djot"

	monkeypatch.setattr(pptx_to_djot.os, "link", lambda _source, _target: (_ for _ in ()).throw(OSError("fail")))
	with pytest.raises(OSError, match="fail"):
		pptx_to_djot.publish_conversion(staging_assets, staging_djot, output_path)

	assert not output_path.exists()
	assert not (tmp_path / "assets" / "deck").exists()


#============================================
def test_pruning_keeps_only_generated_djot_reachable_media(tmp_path: pathlib.Path) -> None:
	"""Only ordinary source images reachable from generated Djot are published."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	for name in ("image_001.png", "shared.png", "figure.png"):
		(staging_assets / name).write_bytes(b"media")
	staging_djot = tmp_path / "deck.djot"
	staging_djot.write_text(
		"=== layout: two-panels\n\n@left\n![Shared](assets/deck/shared.png)\n\n"
		"@right\n![Figure](assets/deck/figure.png)\n",
		encoding="utf-8",
	)

	published_count = pptx_to_djot.prune_staged_media(staging_djot, staging_assets, "deck")
	report = {"unique_media_assets": published_count}

	assert not (staging_assets / "image_001.png").exists()
	assert report["unique_media_assets"] == 2


#============================================
def test_pruning_rejects_unsafe_staged_entries(tmp_path: pathlib.Path) -> None:
	"""Private staging fails closed instead of traversing an unexpected asset entry."""
	staging_assets = tmp_path / "staging-assets"
	staging_assets.mkdir()
	(staging_assets / "shared.png").write_bytes(b"media")
	(staging_assets / "nested").mkdir()
	staging_djot = tmp_path / "deck.djot"
	staging_djot.write_text(
		"=== layout: one-panel\n\n@body\n![Shared](assets/deck/shared.png)\n",
		encoding="utf-8",
	)

	with pytest.raises(ValueError, match="must not contain directories"):
		pptx_to_djot.prune_staged_media(staging_djot, staging_assets, "deck")

	assert (staging_assets / "nested").is_dir()


def test_symlinked_assets_parent_is_rejected_without_clobbering(tmp_path: pathlib.Path) -> None:
	"""A destination assets symlink cannot redirect generated importer output."""
	outside = tmp_path / "outside"
	outside.mkdir()
	(tmp_path / "assets").symlink_to(outside, target_is_directory=True)

	with pytest.raises(ValueError, match="real directory"):
		pptx_to_djot.validate_output_path(tmp_path / "deck.djot")

	assert list(outside.iterdir()) == []


#============================================
def test_invalid_source_geometry_fails_closed() -> None:
	"""Out-of-slide bounds never become permissive normalized geometry."""
	with pytest.raises(ValueError, match="outside slide bounds"):
		geometry.normalized_bounds(-20, 0, 5, 20, 100, 100)

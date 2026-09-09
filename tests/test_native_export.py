"""Public-boundary tests for direct sibling presentation export."""

# Standard Library
import pathlib
import zipfile

# PIP3 Modules
import defusedxml.ElementTree
import pytest
from pptx import Presentation

# Local Modules
import slide_lib.native_export


ODF_NAMESPACES = {
	"draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
	"presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
	"text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


#============================================
def write_deck(tmp_path: pathlib.Path, source: str) -> pathlib.Path:
	"""Write one test-owned inline Djot source."""
	path = tmp_path / "deck.djot"
	path.write_text(source, encoding="utf-8")
	return path


#============================================
def test_direct_renderers_write_editable_sibling_formats(tmp_path: pathlib.Path) -> None:
	"""Both public renderers consume Djot semantics without a conversion bridge."""
	source_path = write_deck(tmp_path,
		"=== layout: one-panel\n\n# Genetics\n\n@body\n\nVisible prose.\n\n"
		"- Parent\n\n  - Child\n")
	deck = slide_lib.native_export.parse_deck(source_path)
	pptx_path = slide_lib.native_export.render_native_pptx(deck, tmp_path / "deck.pptx")
	odp_path = slide_lib.native_export.render_native_odp(deck, tmp_path / "deck.odp")
	assert Presentation(pptx_path).slides[0].shapes
	with zipfile.ZipFile(odp_path) as archive:
		root = defusedxml.ElementTree.fromstring(archive.read("content.xml"))
	page = root.find(".//draw:page", ODF_NAMESPACES)
	assert page is not None and page.attrib[
		f"{{{ODF_NAMESPACES['presentation']}}}presentation-page-layout-name"]
	assert root.find(".//draw:frame[@presentation:class='outline']", ODF_NAMESPACES) is not None
	assert root.find(".//text:list-item/text:list/text:list-item", ODF_NAMESPACES) is not None
	assert "Visible prose." in "".join(root.itertext())


#============================================
def test_folder_discovery_selects_sorted_supported_descendants(tmp_path: pathlib.Path) -> None:
	"""Folder builds retain only Djot decks in stable relative-path order."""
	folder = tmp_path / "slides"
	(folder / "nested").mkdir(parents=True)
	(folder / "z.djot").write_text("=== layout: blank\n", encoding="utf-8")
	(folder / "notes.md").write_text("# Notes\n", encoding="utf-8")
	(folder / "a.djot").write_text("=== layout: blank\n", encoding="utf-8")
	(folder / "nested" / "b.djot").write_text("=== layout: blank\n", encoding="utf-8")
	decks = slide_lib.native_export.discover_decks(str(folder), tmp_path)
	assert [path.relative_to(folder).as_posix() for path in decks] == [
		"a.djot", "nested/b.djot", "z.djot"]


#============================================
def test_input_boundary_rejects_wrong_suffix_and_repository_escape(tmp_path: pathlib.Path) -> None:
	"""Source selection admits only repository-local Djot files."""
	markdown_path = tmp_path / "deck.md"
	markdown_path.write_text("# Not a deck\n", encoding="utf-8")
	with pytest.raises(slide_lib.native_export.PresentationInputError,
			match="use the .djot extension"):
		slide_lib.native_export.validate_input(str(markdown_path), tmp_path)
	repository = tmp_path / "repository"
	repository.mkdir()
	djot_path = write_deck(tmp_path, "=== layout: blank\n")
	with pytest.raises(slide_lib.native_export.PresentationInputError,
			match="inside this repository"):
		slide_lib.native_export.validate_input(str(djot_path), repository)


#============================================
def test_folder_discovery_rejects_empty_source_selection(tmp_path: pathlib.Path) -> None:
	"""A folder without Djot sources receives an actionable input error."""
	folder = tmp_path / "slides"
	folder.mkdir()
	(folder / "notes.md").write_text("# Notes\n", encoding="utf-8")
	with pytest.raises(slide_lib.native_export.PresentationInputError,
			match="no presentation source decks found"):
		slide_lib.native_export.discover_decks(str(folder), tmp_path)
